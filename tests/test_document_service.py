import io
import unittest
from unittest.mock import patch

from app.core.ai_guardrails import verify_file_content
from app.services import document_service
from app.services.document_service import extract_document_structured


class DocumentServiceTests(unittest.TestCase):
    def test_extracts_powerpoint_slide_text_and_preserves_page_markers(self):
        from pptx import Presentation

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Networking Fundamentals"
        slide.placeholders[1].text = (
            "A protocol is a set of rules used to communicate across a network."
        )
        stream = io.BytesIO()
        presentation.save(stream)
        data = stream.getvalue()

        is_valid, message = verify_file_content(data, "lesson.pptx")
        extracted = extract_document_structured(data, "lesson.pptx")

        self.assertTrue(is_valid, message)
        self.assertIn("[Page 1]", extracted["full_text"])
        self.assertIn("Networking Fundamentals", extracted["full_text"])
        self.assertIn("A protocol is a set of rules", extracted["full_text"])
        self.assertEqual(extracted["page_count"], 1)

    def test_pdf_extraction_attempts_ocr_for_image_only_pages(self):
        ocr_text_page = object()

        class FakePage:
            def get_text(self, _kind, **kwargs):
                return (
                    "Recognized text from a scanned page containing several useful terms."
                    if kwargs.get("textpage") is ocr_text_page
                    else ""
                )

            def get_images(self, full=False):
                return [("scanned-image",)] if full else []

            def get_textpage_ocr(self, **kwargs):
                self.ocr_options = kwargs
                return ocr_text_page

        class FakeDocument:
            def __init__(self):
                self.page = FakePage()
                self.closed = False

            def __len__(self):
                return 1

            def load_page(self, _page_num):
                return self.page

            def close(self):
                self.closed = True

        pdf_document = FakeDocument()
        with patch.object(document_service.fitz, "open", return_value=pdf_document):
            extracted = document_service.extract_pdf_structured(b"pdf bytes")

        self.assertEqual(extracted["ocr_pages"], [1])
        self.assertIn("Recognized text", extracted["full_text"])
        self.assertEqual(pdf_document.page.ocr_options["dpi"], 200)
        self.assertTrue(pdf_document.closed)

    def test_scanned_pdf_reports_when_ocr_engine_is_unavailable(self):
        class FakePage:
            def get_text(self, _kind, **_kwargs):
                return ""

            def get_images(self, full=False):
                return [("scanned-image",)] if full else []

            def get_textpage_ocr(self, **_kwargs):
                raise RuntimeError("OCR engine unavailable")

        class FakeDocument:
            def __len__(self):
                return 1

            def load_page(self, _page_num):
                return FakePage()

            def close(self):
                pass

        with patch.object(
            document_service.fitz,
            "open",
            return_value=FakeDocument(),
        ):
            extracted = document_service.extract_pdf_structured(b"pdf bytes")

        self.assertEqual(extracted["full_text"], "")
        self.assertEqual(len(extracted["warnings"]), 1)
        self.assertIn("OCR unavailable for PDF page 1", extracted["warnings"][0])

    def test_text_extraction_normalizes_line_break_hyphenation(self):
        extracted = extract_document_structured(
            b"The net-\nwork transfers information between connected devices.",
            "lesson.txt",
        )

        self.assertIn("network transfers", extracted["full_text"])


if __name__ == "__main__":
    unittest.main()
